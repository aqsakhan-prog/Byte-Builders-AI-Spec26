from pathlib import Path
from math import pi

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from reportlab.lib import colors


OUTPUT_DIR = Path("demo_dataset")
OUTPUT_DIR.mkdir(exist_ok=True)

PART_NUMBER = "2280"
PART_NAME = "CNC Mounting Plate"

# ------------------------------------------------------------
# 1. Generate Excel manufacturing record
# ------------------------------------------------------------

excel_path = OUTPUT_DIR / "manufacturing_records.xlsx"

headers = [
    "Part Number",
    "Part Name",
    "Material",
    "Material Standard",
    "Manufacturing Process",
    "Overall Length (mm)",
    "Overall Width (mm)",
    "Thickness (mm)",
    "Corner Hole Count",
    "Corner Hole Diameter (mm)",
    "Center Hole Diameter (mm)",
    "General Tolerance (mm)",
    "Surface Finish",
    "Quantity",
    "Revision",
    "Drawing File",
    "Model File",
    "Status",
]

record = [
    PART_NUMBER,
    PART_NAME,
    "Aluminium 6061-T6",
    "ASTM B209",
    "CNC milling and drilling",
    100.0,
    60.0,
    10.0,
    4,
    8.0,
    20.0,
    "+/- 0.10",
    "Ra 3.2 micrometres",
    25,
    "A",
    "2280_engineering_drawing.pdf",
    "2280_model.iges",
    "Approved for demonstration",
]

workbook = Workbook()
sheet = workbook.active
sheet.title = "Manufacturing Records"

sheet.append(headers)
sheet.append(record)

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)
thin_border = Border(
    left=Side(style="thin", color="B7B7B7"),
    right=Side(style="thin", color="B7B7B7"),
    top=Side(style="thin", color="B7B7B7"),
    bottom=Side(style="thin", color="B7B7B7"),
)

for cell in sheet[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

for row in sheet.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

sheet.freeze_panes = "A2"
sheet.auto_filter.ref = sheet.dimensions
sheet.row_dimensions[1].height = 45
sheet.row_dimensions[2].height = 35

for column_cells in sheet.columns:
    maximum_length = max(
        len(str(cell.value)) if cell.value is not None else 0
        for cell in column_cells
    )
    column_letter = column_cells[0].column_letter
    sheet.column_dimensions[column_letter].width = min(maximum_length + 3, 30)

workbook.save(excel_path)

# ------------------------------------------------------------
# 2. Generate PDF engineering drawing
# ------------------------------------------------------------

pdf_path = OUTPUT_DIR / "2280_engineering_drawing.pdf"
page_width, page_height = landscape(A4)

drawing = canvas.Canvas(str(pdf_path), pagesize=landscape(A4))
drawing.setTitle("Part 2280 Engineering Drawing")
drawing.setAuthor("SPEC'26 Hackathon Team - Demonstration Dataset")
drawing.setSubject("Engineering drawing for CNC Mounting Plate")

# Border
drawing.setStrokeColor(colors.black)
drawing.setLineWidth(1.2)
drawing.rect(20, 20, page_width - 40, page_height - 40)

# Title
drawing.setFont("Helvetica-Bold", 18)
drawing.drawString(40, page_height - 50, "ENGINEERING DRAWING - CNC MOUNTING PLATE")

drawing.setFont("Helvetica", 10)
drawing.drawRightString(
    page_width - 40,
    page_height - 48,
    "Team-created demonstration dataset",
)

# Drawing scale: 4 points for each millimetre
scale = 4.0
origin_x = 120
origin_y = 210
plate_width = 100 * scale
plate_height = 60 * scale

# Plate outline
drawing.setLineWidth(2)
drawing.rect(origin_x, origin_y, plate_width, plate_height)

# Center lines
drawing.setDash(4, 3)
drawing.setStrokeColor(colors.grey)
drawing.line(
    origin_x + plate_width / 2,
    origin_y - 20,
    origin_x + plate_width / 2,
    origin_y + plate_height + 20,
)
drawing.line(
    origin_x - 20,
    origin_y + plate_height / 2,
    origin_x + plate_width + 20,
    origin_y + plate_height / 2,
)
drawing.setDash()
drawing.setStrokeColor(colors.black)

# Holes: four DIA 8 holes and one DIA 20 hole
corner_holes = [(10, 10), (90, 10), (10, 50), (90, 50)]

for x_mm, y_mm in corner_holes:
    x = origin_x + x_mm * scale
    y = origin_y + y_mm * scale
    drawing.circle(x, y, 4 * scale, stroke=1, fill=0)
    drawing.line(x - 22, y, x + 22, y)
    drawing.line(x, y - 22, x, y + 22)

center_x = origin_x + 50 * scale
center_y = origin_y + 30 * scale
drawing.circle(center_x, center_y, 10 * scale, stroke=1, fill=0)
drawing.line(center_x - 48, center_y, center_x + 48, center_y)
drawing.line(center_x, center_y - 48, center_x, center_y + 48)

# Overall length dimension
dimension_y = origin_y - 55
drawing.setLineWidth(1)
drawing.line(origin_x, origin_y, origin_x, dimension_y)
drawing.line(origin_x + plate_width, origin_y, origin_x + plate_width, dimension_y)
drawing.line(origin_x, dimension_y, origin_x + plate_width, dimension_y)
drawing.line(origin_x, dimension_y - 5, origin_x + 10, dimension_y + 5)
drawing.line(origin_x + plate_width, dimension_y - 5, origin_x + plate_width - 10, dimension_y + 5)
drawing.setFont("Helvetica-Bold", 11)
drawing.drawCentredString(origin_x + plate_width / 2, dimension_y + 8, "100 mm")

# Overall width dimension
dimension_x = origin_x - 55
drawing.line(origin_x, origin_y, dimension_x, origin_y)
drawing.line(origin_x, origin_y + plate_height, dimension_x, origin_y + plate_height)
drawing.line(dimension_x, origin_y, dimension_x, origin_y + plate_height)
drawing.line(dimension_x - 5, origin_y, dimension_x + 5, origin_y + 10)
drawing.line(
    dimension_x - 5,
    origin_y + plate_height,
    dimension_x + 5,
    origin_y + plate_height - 10,
)
drawing.saveState()
drawing.translate(dimension_x - 10, origin_y + plate_height / 2)
drawing.rotate(90)
drawing.drawCentredString(0, 0, "60 mm")
drawing.restoreState()

# Hole callouts
drawing.setFont("Helvetica-Bold", 11)
drawing.drawString(origin_x + plate_width + 45, origin_y + plate_height - 20, "HOLE CALLOUTS")
drawing.setFont("Helvetica", 10)
drawing.drawString(origin_x + plate_width + 45, origin_y + plate_height - 40, "4 x DIA 8 mm THRU")
drawing.drawString(origin_x + plate_width + 45, origin_y + plate_height - 58, "Corner centres: 10 mm")
drawing.drawString(origin_x + plate_width + 45, origin_y + plate_height - 74, "from adjacent edges")
drawing.drawString(origin_x + plate_width + 45, origin_y + plate_height - 100, "1 x DIA 20 mm THRU")
drawing.drawString(origin_x + plate_width + 45, origin_y + plate_height - 118, "Located at plate centre")

# Notes box
notes_x = 40
notes_y = 45
notes_width = 470
notes_height = 110

drawing.rect(notes_x, notes_y, notes_width, notes_height)
drawing.setFont("Helvetica-Bold", 11)
drawing.drawString(notes_x + 10, notes_y + notes_height - 18, "MANUFACTURING NOTES")

notes = [
    "1. Material: Aluminium 6061-T6, ASTM B209.",
    "2. Overall dimensions: 100 x 60 x 10 mm.",
    "3. General dimensional tolerance: +/- 0.10 mm unless otherwise specified.",
    "4. Surface finish: Ra 3.2 micrometres.",
    "5. Process: CNC milling and drilling.",
    "6. Deburr and remove all sharp edges.",
]

drawing.setFont("Helvetica", 9)
text_y = notes_y + notes_height - 36

for note in notes:
    drawing.drawString(notes_x + 10, text_y, note)
    text_y -= 13

# Title block
block_x = 530
block_y = 45
block_width = page_width - block_x - 40
block_height = 110

drawing.rect(block_x, block_y, block_width, block_height)

title_rows = [
    ("PART NUMBER", PART_NUMBER),
    ("PART NAME", PART_NAME),
    ("REVISION", "A"),
    ("SCALE", "NTS"),
    ("UNITS", "millimetres"),
    ("SHEET", "1 of 1"),
]

row_height = block_height / len(title_rows)

for index, (label, value) in enumerate(title_rows):
    row_top = block_y + block_height - index * row_height
    if index > 0:
        drawing.line(block_x, row_top, block_x + block_width, row_top)

    drawing.setFont("Helvetica-Bold", 8)
    drawing.drawString(block_x + 7, row_top - 13, label)

    drawing.setFont("Helvetica", 9)
    drawing.drawString(block_x + 105, row_top - 13, value)

drawing.setFont("Helvetica-Oblique", 8)
drawing.drawRightString(
    page_width - 40,
    27,
    "Created for SPEC'26 CNC Part Analyzer demonstration",
)

drawing.save()

# ------------------------------------------------------------
# Completion report
# ------------------------------------------------------------

gross_volume = 100 * 60 * 10
hole_volume = (4 * pi * (4 ** 2) * 10) + (pi * (10 ** 2) * 10)
estimated_net_volume = gross_volume - hole_volume

print()
print("Demo dataset files generated successfully:")
print(f"1. {excel_path.resolve()}")
print(f"2. {pdf_path.resolve()}")
print()
print(f"Expected model bounding box: 100 x 60 x 10 mm")
print(f"Expected net model volume: {estimated_net_volume:.2f} cubic mm")